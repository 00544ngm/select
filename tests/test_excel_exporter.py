from __future__ import annotations

from openpyxl import load_workbook

from app.domain.dto import (
    DirectionDTO,
    HypothesisDTO,
    HypothesisResultDTO,
    JudgmentResultDTO,
)
from app.infrastructure.storage.excel_exporter import (
    export_hypothesis_to_excel,
    export_judgment_to_excel,
)


def test_empty_judgment_exports(tmp_path):
    path = export_judgment_to_excel(
        JudgmentResultDTO(),
        tmp_path / "empty.xlsx",
    )

    assert path.exists()
    workbook = load_workbook(path, read_only=True)
    assert "判级概览" in workbook.sheetnames


def test_alignment_product_is_included_without_score_data(tmp_path):
    result = JudgmentResultDTO(
        alignment_review=[{"product_b": "Basket", "overall_verdict": "alignment"}],
        final_grade="B",
        priority_score=0.5,
    )

    path = export_judgment_to_excel(result, tmp_path / "alignment.xlsx")

    workbook = load_workbook(path, read_only=True, data_only=True)
    overview = workbook["判级概览"]
    values = [cell.value for row in overview.iter_rows() for cell in row]
    assert "Basket" in values


def test_hypothesis_excel_exports_v2_score_breakdown_and_rejection(tmp_path):
    result = HypothesisResultDTO(
        directions=[
            DirectionDTO(hypothesis=HypothesisDTO(
                direction_name="Matching Ink",
                model_version="combination_model_v2.0",
                final_score=89,
                raw_score=91,
                score_cap=89,
                recommendation_level="focus",
                primary_relation="spec_compatibility",
                consumer_simulation="A",
                score_breakdown={
                    "relation_strength": 5,
                    "lifecycle_connection": 5,
                    "market_evidence": 8,
                },
            ))
        ]
    )

    path = export_hypothesis_to_excel(result, tmp_path / "hypothesis-v2.xlsx")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["辅品方向"]
    headers = [cell.value for cell in next(sheet.iter_rows())]
    values = list(next(sheet.iter_rows(min_row=2)))
    row = dict(zip(headers, [cell.value for cell in values]))

    assert "最终评分" in headers
    assert "原始计算分" in headers
    assert "分数上限" in headers
    assert "淘汰原因" in headers
    assert row["最终评分"] == 89
    assert row["原始计算分"] == 91
    assert row["分数上限"] == 89
    assert row["关系强度"] == 5


def test_zero_direction_v21_excel_has_reliability_summary_and_status_row(tmp_path):
    result = HypothesisResultDTO(
        model_version="combination_model_v2.1",
        result_status="completed_no_qualified_candidates",
        result_message="分析已完成，未发现达到高粘性门槛的辅品，不建议为了凑数量强行组合。",
        raw_direction_count=0,
        qualified_direction_count=0,
        hold_direction_count=0,
        rejected_direction_count=0,
        rejection_summary={"food_blocked": 2},
        audit_performed=True,
        audit_reason="initial_v2.1_directions_empty",
        initial_raw_direction_count=0,
        audit_raw_direction_count=0,
        audit_outcome="confirmed_no_candidates",
        provider="custom",
        provider_model="claude-fable-5",
    )

    path = export_hypothesis_to_excel(result, tmp_path / "v21-zero.xlsx")
    workbook = load_workbook(path, read_only=True, data_only=True)

    assert "结果摘要" in workbook.sheetnames
    summary_values = [
        cell.value for row in workbook["结果摘要"].iter_rows() for cell in row
    ]
    assert "combination_model_v2.1" in summary_values
    assert "custom" in summary_values
    assert "claude-fable-5" in summary_values
    assert "不可直接采购" in summary_values
    assert "food_blocked: 2" in summary_values

    directions = workbook["辅品方向"]
    headers = [cell.value for cell in next(directions.iter_rows())]
    values = [cell.value for cell in next(directions.iter_rows(min_row=2))]
    row = dict(zip(headers, values))
    assert row["方向名称"] == "未发现合格候选"
    assert row["执行状态"] == "completed_no_qualified_candidates"
    assert row["英文关键词"] is None
    assert row["最终评分"] is None


def test_excel_requires_qualified_top_level_status_for_actionable_result(tmp_path):
    result = HypothesisResultDTO(
        model_version="combination_model_v2.1",
        result_status="completed_needs_evidence",
        directions=[
            DirectionDTO(
                hypothesis=HypothesisDTO(
                    direction_name="Contradictory pass candidate",
                    execution_status="pass",
                    decision_action="priority_test",
                )
            )
        ],
    )

    path = export_hypothesis_to_excel(result, tmp_path / "contradictory-status.xlsx")
    workbook = load_workbook(path, read_only=True, data_only=True)
    summary = {
        row[0].value: row[1].value
        for row in workbook["结果摘要"].iter_rows(min_row=2)
    }

    assert summary["是否可用于采购决策"] == "不可直接采购"


def test_excel_uses_employee_action_language_for_hold_and_reject(tmp_path):
    result = HypothesisResultDTO(
        model_version="combination_model_v2.1",
        directions=[
            DirectionDTO(
                hypothesis=HypothesisDTO(
                    direction_name="待核对笔套装",
                    execution_status="hold",
                    missing_evidence=["核对笔夹与笔杆直径"],
                )
            ),
            DirectionDTO(
                hypothesis=HypothesisDTO(
                    direction_name="已确认不兼容笔套装",
                    execution_status="reject",
                    compatibility_status="blocked",
                    incompatibility_reason="主品孔径小于候选直径",
                    source_fact_ids=["spec:main", "spec:candidate"],
                    rejection_codes=["incompatible"],
                )
            ),
        ],
    )

    path = export_hypothesis_to_excel(result, tmp_path / "employee-guidance.xlsx")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["辅品方向"]
    headers = [cell.value for cell in next(sheet.iter_rows())]
    rows = [
        dict(zip(headers, [cell.value for cell in row]))
        for row in sheet.iter_rows(min_row=2)
    ]

    assert rows[0]["当前处理建议"] == "补充证据后复核"
    assert rows[0]["下一步"] == "请核对：核对笔夹与笔杆直径"
    assert rows[1]["当前处理建议"] == "当前方案暂不进入测试"
    assert rows[1]["具体原因"] == "主品孔径小于候选直径"
